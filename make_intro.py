# -*- coding: utf-8 -*-
"""生成《个人简介.docx》（招新材料），格式：黑体标题 + 宋体正文。"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "docs", "个人简介（招新材料）.docx")


def set_run(run, cn="宋体", size=12, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), cn)


def para(doc, text, cn="宋体", size=12, bold=False, align=None, indent=True, after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    p.paragraph_format.space_after = Pt(after)
    set_run(p.add_run(text), cn=cn, size=size, bold=bold)
    return p


doc = Document()
for s in doc.sections:
    s.top_margin = Cm(2.5)
    s.bottom_margin = Cm(2.5)
    s.left_margin = Cm(2.8)
    s.right_margin = Cm(2.8)

para(doc, "个 人 简 介", cn="黑体", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=12)
para(doc, "〔姓名〕　〔学号〕　〔年级班级〕　光电信息科学与工程专业", size=12,
     align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=14)

para(doc, "一、个人概述", cn="黑体", size=14, bold=True, indent=False)
para(doc, "我是深圳大学光电信息科学与工程专业〔20XX〕级学生〔姓名〕。作为一名理工科学生，"
          "我始终相信“实践出真知”：课堂上学到的原理，只有亲手做成看得见、摸得着的东西，"
          "才算真正学会。光电实验、金工实习和自主项目开发让我对“从想法到作品”的完整过程充满热情，"
          "也养成了遇到问题先分析原因、再动手解决的做事习惯。")

para(doc, "二、专业能力与动手实践", cn="黑体", size=14, bold=True, indent=False)
para(doc, "· 专业基础：系统学习大学物理、高等数学、电路等课程，对光学、光电探测等方向有扎实的理论基础；")
para(doc, "· 动手实践：完成金工实习，熟悉常用加工工艺与工具操作；长期在课程实验与自主项目中练习数据采集、"
          "分析与处理，能够独立完成从实验设计到结果验证的完整流程；")
para(doc, "· 项目开发：独立开发《光电实验数据处理助手》（物光创新实验室 2026 招新作品），"
          "将光学公式计算、实验数据拟合、不确定度评定与实验报告生成整合为一个本地工具。"
          "项目采用“算法与界面分层”的模块化设计，自定义函数拟合算法（Levenberg-Marquardt）自主实现，"
          "代码全部开源、数据可复现，并在开发过程中主动解决过数值算法不收敛等问题。")

para(doc, "三、公益与实践经历", cn="黑体", size=14, bold=True, indent=False)
para(doc, "· 2024 年暑期，作为“追光童行”支教突击队员，赴龙门县白芒坑革命旧址群开展乡村支教，"
          "把知识带给乡村的孩子，也在真实环境中锻炼了沟通表达与团队协作能力；")
para(doc, "· 参与深圳本地调研，走进一线了解社会，积累了实地观察、数据收集与报告撰写的经验。")

para(doc, "四、加入实验室的动机", cn="黑体", size=14, bold=True, indent=False)
para(doc, "我希望加入物光创新实验室，把光电专业知识与工程实践结合起来，在真实项目中学习"
          "从选题、实现、验证到展示的完整能力；同时愿意承担实验室的日常任务，与大家互相学习、共同进步。")

para(doc, "五、联系方式", cn="黑体", size=14, bold=True, indent=False)
para(doc, "邮箱：〔邮箱〕　　电话：〔电话〕　　GitHub：〔GitHub 用户名〕", indent=False)

doc.save(OUT)
print("已生成：", OUT)
