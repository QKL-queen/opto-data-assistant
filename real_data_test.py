# -*- coding: utf-8 -*-
"""真实实验数据实测脚本

用网络上公开的真实实验数据，逐一测试项目的全部功能（模拟界面点击路径）：

数据来源（均为公开实验教材/报告，2026-10-12 检索）：
- 杨氏双缝干涉（Virtual Labs, 钠黄光 589nm 5 组实测）：
  https://virtual-labs.github.io/exp-wave-nature-of-light-dei/procedure.html
  D/m ∈ {1.0, 1.5, 2.0, 1.0, 1.5}，d/mm ∈ {0.50, 0.50, 0.50, 1.00, 1.00}
  β/mm ∈ {1.178, 1.767, 2.356, 0.589, 0.884}
- 光电效应测普朗克常数（中国科学技术大学实验报告，汞灯 5 谱线零电流法）：
  http://home.ustc.edu.cn/~wuyuewy/dwsy/report/guangdian.pdf
  λ/nm = 365.0, 404.7, 435.8, 546.1, 577.0
  ν/×10^14Hz = 8.214, 7.408, 6.879, 5.490, 5.196
  U0/V = 1.704, 1.304, 1.086, 0.534, 0.436 （报告拟合 h=6.645×10^-34 J·s）
- 光栅 600 lines/mm + He-Ne 632.8nm（William & Mary 物理实验手册）：
  http://physics.wm.edu/~labs/108_pdf/ch11.pdf
- 单缝衍射 He-Ne 632.8nm 缝宽 0.08mm（上海交大物理实验中心教材同款配置）：
  https://pec.sjtu.edu.cn/ocw/materials/b/T09.pdf
- 双缝干涉测波长（北京恒奥德 J2515 仪器实测，红色滤色片 4 次重复读数）：
  http://www.bjhadkj.com/bjhadkj-Products-24409403/
  X1: 3.82,3.80,3.80,3.80 mm；X7: 14.74,14.76,14.76,14.76 mm；d=0.250mm, L=700mm, n=7
  λ = d·(X7-X1)/(L·6) = 652.4 nm（标称 655±10nm）

运行：python real_data_test.py
"""

import math
import os
import re
import sys
import tempfile

import app  # 界面层（直接调用与点击按钮相同的处理函数）

PASS = 0
FAIL = 0
FAILED = []


def check(name, ok, detail=""):
    global PASS, FAIL
    if ok:
        PASS += 1
        print(f"  [PASS] {name}")
    else:
        FAIL += 1
        FAILED.append(name)
        print(f"  [FAIL] {name}  {detail}")


def get_num(md_text, key_cn):
    """从界面返回的 Markdown 中按中文物理量名提取数值。"""
    pat = re.compile(rf"-\s*\*\*{re.escape(key_cn)}\*\*\s*=\s*([0-9.eE+-]+)")
    m = pat.search(md_text)
    if not m:
        return None
    return float(m.group(1))


def get_line_value(md_text, prefix):
    """按行前缀提取数值（如「斜率 k = ...」）。"""
    for line in md_text.splitlines():
        if line.startswith(prefix):
            m = re.search(r"=\s*([0-9.eE+-]+)", line)
            if m:
                return float(m.group(1))
    return None


def get_inline(md_text, label):
    """从整段文本中匹配 label = 数值（label 可含或不含 =，如 "a =" 或 "平均值 x̄"）。"""
    label2 = label.rstrip().rstrip("=").rstrip()
    m = re.search(re.escape(label2) + r"\s*=\s*([0-9.eE+-]+)", md_text)
    return float(m.group(1)) if m else None


# ============================================================
# 1. 公式计算器（14 公式之一部分，用真实数据）
# ============================================================
print("=" * 50)
print("[1/6] 公式计算器（真实实验数据）")
print("=" * 50)

# --- 1a. 杨氏双缝干涉：Virtual Labs 表 1（钠黄光 589nm）---
# Δx = λL/d；正求条纹间距与实测 β 比对
young_rows = [
    # (L/m, d/mm, β/mm 实测)
    (1.0, 0.50, 1.178),
    (1.5, 0.50, 1.767),
    (2.0, 0.50, 2.356),
    (1.0, 1.00, 0.589),
    (1.5, 1.00, 0.884),
]
for i, (L, d, beta) in enumerate(young_rows, 1):
    txt = app.compute_formula("杨氏双缝干涉条纹间距", 589.0, L, d, None)
    dx = get_num(txt, "条纹间距 Δx")
    ok = dx is not None and abs(dx - beta) / beta < 0.01
    check(f"杨氏双缝 第{i}组 Δx≈{beta}mm", ok, f"计算得 Δx={dx}mm")
# 反求波长：第 1 组 Δx=1.178, L=1.0, d=0.5 → λ ≈ 589nm
txt = app.compute_formula("杨氏双缝干涉条纹间距", None, 1.0, 0.5, 1.178)
lam = get_num(txt, "波长 λ")
check("杨氏双缝 反求波长 λ≈589nm", lam is not None and abs(lam - 589) / 589 < 0.01, f"λ={lam}nm")

# --- 1b. 光栅方程：600 lines/mm + He-Ne 632.8nm，一级衍射角 ---
# d = 1/600 mm = 1.6667μm；θ1 = asin(λ/d) = asin(632.8e-9/1.6667e-6) ≈ 22.32°
txt = app.compute_formula("光栅方程", 1.6667, None, 1, 632.8)
theta = get_num(txt, "衍射角 θ")
expect = math.degrees(math.asin(632.8e-9 / (1.6667e-6)))
check("光栅方程 θ1≈22.3°（600线/mm+HeNe）", theta is not None and abs(theta - expect) < 0.01,
      f"θ={theta}°，理论 {expect:.3f}°")
# 反求光栅常数
txt = app.compute_formula("光栅方程", None, 22.32, 1, 632.8)
d = get_num(txt, "光栅常数 d")
check("光栅方程 反求 d≈1.667μm", d is not None and abs(d - 1.6667) / 1.6667 < 0.01, f"d={d}μm")

# --- 1c. 单缝衍射：He-Ne 632.8nm、缝宽 0.08mm，一级暗纹 ---
# θ1 = asin(λ/a) = asin(632.8e-9/0.08e-3) ≈ 0.453°
txt = app.compute_formula("单缝衍射暗纹条件", 0.08, None, 1, 632.8)
theta = get_num(txt, "衍射角 θ")
expect = math.degrees(math.asin(632.8e-9 / (0.08e-3)))
check("单缝衍射 θ1≈0.453°（HeNe+0.08mm）", theta is not None and abs(theta - expect) < 0.01,
      f"θ={theta}°，理论 {expect:.3f}°")
# 反求缝宽
txt = app.compute_formula("单缝衍射暗纹条件", None, 0.4532, 1, 632.8)
a = get_num(txt, "缝宽 a")
check("单缝衍射 反求 a≈0.08mm", a is not None and abs(a - 0.08) / 0.08 < 0.01, f"a={a}mm")

# --- 1d. 波长-频率关系：中科大汞灯谱线 365.0nm → 8.214×10^14Hz = 821.4THz ---
txt = app.compute_formula("波长-频率关系", 365.0, None)
nu = get_num(txt, "频率 ν")
check("波长-频率 365nm→821.4THz", nu is not None and abs(nu - 821.4) / 821.4 < 0.01, f"ν={nu}THz")

# --- 1e. 光子能量：汞灯 577.0nm → 2.149eV ---
txt = app.compute_formula("光子能量", 577.0, None)
E = get_num(txt, "光子能量 E")
check("光子能量 577nm→2.15eV", E is not None and abs(E - 2.149) / 2.149 < 0.01, f"E={E}eV")

# --- 1f. 薄透镜成像：标准例题 u=20, v=30 → f=12cm ---
txt = app.compute_formula("薄透镜成像公式", None, 20.0, 30.0)
f = get_num(txt, "焦距 f")
check("薄透镜 u=20,v=30→f=12cm", f is not None and abs(f - 12.0) < 1e-6, f"f={f}cm")

# ============================================================
# 2. 数据拟合（真实数据）
# ============================================================
print("=" * 50)
print("[2/6] 数据拟合（真实数据）")
print("=" * 50)

# --- 2a. 光电效应 U0-ν 线性拟合 → 斜率 k=h/e，h≈6.64×10^-34 ---
pe_data = "8.214, 1.704\n7.408, 1.304\n6.879, 1.086\n5.490, 0.534\n5.196, 0.436"
txt, fig, path = app.do_fit(pe_data, "linear", 2, "", "", "频率 ν / 10^14 Hz", "遏止电压 U0 / V")
k = get_line_value(txt, "斜率 k")
h_est = k * 1.602e-33 if k else None  # ν 以 10^14 Hz 为单位：h = k·10^-14·e
check("光电效应 拟合斜率≈0.414", k is not None and abs(k - 0.4143) / 0.4143 < 0.02, f"k={k}")
check("光电效应 拟合 h≈6.64e-34（±5%）",
      h_est is not None and abs(h_est - 6.626e-34) / 6.626e-34 < 0.05, f"h={h_est:.3g}")
r2 = get_line_value(txt, "决定系数 R²")
check("光电效应 拟合 R²>0.99", r2 is not None and r2 > 0.99, f"R²={r2}")
check("光电效应 拟合图已生成", fig is not None and os.path.exists(path))

# --- 2b. 杨氏双缝 D-β 线性拟合（d=0.5mm 前三组）→ 斜率=λ/d → λ≈589nm ---
young_fit = "1.0, 1.178\n1.5, 1.767\n2.0, 2.356"
txt, fig, path = app.do_fit(young_fit, "linear", 2, "", "", "缝屏距离 L / m", "条纹间距 β / mm")
k = get_line_value(txt, "斜率 k")
lam_est = k * 0.5 if k else None  # λ[nm] = k[mm/m] × d[mm] × 1e-3? 需换算
# β[mm] = λ[nm]·L[m]/d[mm] → 斜率 k[mm/m] = λ·1e-3/d → λ[nm] = k·d·1e3
lam_est = k * 0.5 * 1e3 if k else None
check("杨氏双缝 D-β 拟合 λ≈589nm（±3%）",
      lam_est is not None and abs(lam_est - 589) / 589 < 0.03, f"λ={lam_est}nm, k={k}")

# --- 2c. 自定义拟合：光电效应用 y = a*x + b 的形式？改用指数衰减验证自定义拟合 ---
# 用朗伯-比尔型衰减数据：I = 1.0·exp(-0.5·x)，加轻微真实噪声
lb_data = "0, 1.00\n1, 0.61\n2, 0.37\n3, 0.22\n4, 0.14\n5, 0.08"
txt, fig, path = app.do_fit(lb_data, "custom", 2, "a*exp(-b*x)+c", "1, 0.5, 0.01",
                            "光程 x / cm", "光强 I / 相对值")
a_param = get_inline(txt, "a =")
b_param = get_inline(txt, "b =")
check("自定义拟合 a≈1.0（朗伯-比尔）", a_param is not None and abs(a_param - 1.0) < 0.15, f"a={a_param}")
check("自定义拟合 b≈0.5（吸收系数）", b_param is not None and abs(b_param - 0.5) < 0.15, f"b={b_param}")

# --- 2d. CSV 导入（把真实光电效应数据写成 CSV 再读入） ---
csv_path = os.path.join(tempfile.gettempdir(), "真实数据_光电效应.csv")
with open(csv_path, "w", encoding="utf-8-sig") as fp:
    fp.write("频率, 遏止电压\n" + pe_data)
csv_txt = app.load_file_to_text(csv_path)
check("CSV 导入读取内容", "8.214, 1.704" in csv_txt, csv_txt[:40])
txt, fig, path = app.do_fit(csv_txt, "linear", 2, "", "", "ν", "U0")
check("CSV 导入后可直接拟合", "斜率 k" in txt, txt[:50])
os.remove(csv_path)

# ============================================================
# 3. 不确定度评定（真实重复读数）
# ============================================================
print("=" * 50)
print("[3/6] 不确定度评定（真实重复读数）")
print("=" * 50)

# 北京恒奥德 J2515 双缝干涉仪：测微目镜游标读数 X1 重复 4 次，游标最小读数 0.02mm
txt = app.do_direct_uncertainty("3.82\n3.80\n3.80\n3.80", 0.02, True, 2)
mean = get_inline(txt, "平均值 x̄")
check("直接测量 平均值≈3.805mm", mean is not None and abs(mean - 3.805) < 0.01, f"x̄={mean}")
check("直接测量 含扩展不确定度", "U =" in txt and "±" in txt)
check("直接测量 结果式完整", "x = x̄ ± U" in txt)

# ============================================================
# 4. 误差传递（真实公式：双缝干涉测波长）
# ============================================================
print("=" * 50)
print("[4/6] 误差传递（双缝干涉测波长 λ = d·(X7-X1)/(L·6)）")
print("=" * 50)
txt = app.do_propagation("(x7-x1)*d/(L*6)",
                         "x7 = 14.76, 0.01\nx1 = 3.80, 0.01\nd = 0.250, 0.003\nL = 700, 2")
y = get_inline(txt, "计算结果：y")
check("误差传递 λ≈652.4nm", y is not None and abs(y - 0.0006524) / 0.0006524 < 0.01, f"y={y}mm（={y*1e6:.0f}nm）")
check("误差传递 含合成不确定度", "不确定度" in txt or "uc" in txt or "Uc" in txt or "u(y)" in txt)

# ============================================================
# 5. 报告生成（光电效应真实数据，线性拟合 → docx 回读）
# ============================================================
print("=" * 50)
print("[5/6] 报告生成（光电效应真实数据）")
print("=" * 50)
preview, out_path = app.gen_report(
    "光电效应测普朗克常数", "王子骞", "2025270487", "2025级光电01班",
    "2026-10-12", "（指导教师）",
    "验证爱因斯坦光电效应方程，测量普朗克常数。",
    "遏止电压与入射光频率成线性关系 U0 = (h/e)·ν - W/e。",
    "汞灯、光电管、滤色片（365.0~577.0nm）、直流电压源、微电流计。",
    "",  # 结论留空 → 自动生成
    pe_data, "linear", 2, "", "",
    "频率 ν / 10^14 Hz", "遏止电压 U0 / V",
    "", 0.0,
)
check("报告 生成预览", preview is not None and "报告预览" in preview, str(preview)[:60])
check("报告 预览含中文拟合方式", "线性拟合" in preview, str(preview)[:80])
check("报告 docx 已生成", out_path is not None and os.path.exists(out_path), out_path)
if out_path and os.path.exists(out_path):
    import docx
    d = docx.Document(out_path)
    texts = "\n".join(p.text for p in d.paragraphs)
    tables = len(d.tables)
    check("报告 含实验名称", "光电效应测普朗克常数" in texts)
    check("报告 含数据处理章节", "五、数据处理" in texts or "数据处理" in texts)
    check("报告 含结论章节", "七、结论" in texts or "结论" in texts)
    check("报告 含拟合结果文本", "斜率" in texts or "R²" in texts)
    check("报告 含数据表格", tables >= 1, f"表格数={tables}")

# ============================================================
# 6. 边界与健壮性（沿用真实数据口径）
# ============================================================
print("=" * 50)
print("[6/6] 边界与健壮性")
print("=" * 50)
# 缺数据拦截
txt, _, _ = app.do_fit("", "linear", 2, "", "", "x", "y")
check("空数据被拦截", "请先粘贴数据" in txt, repr(txt))
# 数据点不足
txt, _, _ = app.do_fit("1, 2", "linear", 2, "", "", "x", "y")
check("数据点不足被拦截", "不足" in txt or "至少" in txt, repr(txt))
# 公式缺两个输入
txt = app.compute_formula("薄透镜成像公式", None, None, 30.0)
check("公式留空两项被拦截", "一次只能求一个" in txt or "至少给出" in txt)
# 全反射物理约束
txt = app.compute_formula("全反射临界角", None, 1.0, 1.5)
check("全反射 n1<n2 被拦截", "n1 > n2" in txt or "光密" in txt)

print("=" * 50)
print(f"结果：通过 {PASS} 项，失败 {FAIL} 项")
if FAILED:
    print("失败项：")
    for f in FAILED:
        print("  -", f)
print("=" * 50)
sys.exit(1 if FAIL else 0)
